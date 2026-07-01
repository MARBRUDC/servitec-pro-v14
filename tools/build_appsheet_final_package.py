from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "appsheet_digemid_final"
XLSX = OUT_DIR / "DIGEMID_Drogueria_AppSheet.xlsx"
IMPLEMENTACION = OUT_DIR / "IMPLEMENTACION_FINAL.md"


TABLES = {
    "CONFIG": [
        "ID_CONFIG",
        "RAZON_SOCIAL",
        "RUC",
        "DIRECCION",
        "TELEFONO",
        "EMAIL",
        "RESPONSABLE_TECNICO",
        "DIRECTOR_TECNICO",
        "LOGO",
        "DIAS_ALERTA",
    ],
    "PROVEEDORES": [
        "ID_PROVEEDOR",
        "RUC",
        "RAZON_SOCIAL",
        "DIRECCION",
        "TELEFONO",
        "EMAIL",
        "CONTACTO",
        "ESTADO",
    ],
    "CLIENTES": [
        "ID_CLIENTE",
        "RUC_DNI",
        "NOMBRE",
        "DIRECCION",
        "TELEFONO",
        "EMAIL",
        "ESTADO",
    ],
    "PRODUCTOS": [
        "ID_PRODUCTO",
        "CODIGO",
        "NOMBRE",
        "MARCA",
        "LABORATORIO",
        "REGISTRO_SANITARIO",
        "PRESENTACION",
        "UNIDAD",
        "TEMPERATURA",
        "ESTADO",
    ],
    "LOTES": [
        "ID_LOTE",
        "ID_PRODUCTO",
        "LOTE",
        "FECHA_FABRICACION",
        "FECHA_VENCIMIENTO",
        "PRECIO_COMPRA",
        "PRECIO_VENTA",
    ],
    "INGRESOS": [
        "ID_INGRESO",
        "FECHA",
        "ID_PROVEEDOR",
        "ID_PRODUCTO",
        "ID_LOTE",
        "GUIA",
        "FACTURA",
        "CANTIDAD",
        "USUARIO",
    ],
    "EGRESOS": [
        "ID_EGRESO",
        "FECHA",
        "ID_CLIENTE",
        "ID_PRODUCTO",
        "ID_LOTE",
        "GUIA",
        "FACTURA",
        "CANTIDAD",
        "USUARIO",
    ],
    "DOCUMENTOS": [
        "ID_DOCUMENTO",
        "ID_LOTE",
        "TIPO_DOCUMENTO",
        "ARCHIVO",
        "OBSERVACION",
    ],
}

KEYS = {
    "CONFIG": "ID_CONFIG",
    "PROVEEDORES": "ID_PROVEEDOR",
    "CLIENTES": "ID_CLIENTE",
    "PRODUCTOS": "ID_PRODUCTO",
    "LOTES": "ID_LOTE",
    "INGRESOS": "ID_INGRESO",
    "EGRESOS": "ID_EGRESO",
    "DOCUMENTOS": "ID_DOCUMENTO",
}

LABELS = {
    "CONFIG": "RAZON_SOCIAL",
    "PROVEEDORES": "RAZON_SOCIAL",
    "CLIENTES": "NOMBRE",
    "PRODUCTOS": "NOMBRE",
    "LOTES": "LOTE",
    "INGRESOS": "GUIA",
    "EGRESOS": "GUIA",
    "DOCUMENTOS": "TIPO_DOCUMENTO",
}

REFS = {
    ("LOTES", "ID_PRODUCTO"): "PRODUCTOS",
    ("INGRESOS", "ID_PROVEEDOR"): "PROVEEDORES",
    ("INGRESOS", "ID_PRODUCTO"): "PRODUCTOS",
    ("INGRESOS", "ID_LOTE"): "LOTES",
    ("EGRESOS", "ID_CLIENTE"): "CLIENTES",
    ("EGRESOS", "ID_PRODUCTO"): "PRODUCTOS",
    ("EGRESOS", "ID_LOTE"): "LOTES",
    ("DOCUMENTOS", "ID_LOTE"): "LOTES",
}

ENUMS = {
    ("PROVEEDORES", "ESTADO"): ["ACTIVO", "INACTIVO"],
    ("CLIENTES", "ESTADO"): ["ACTIVO", "INACTIVO"],
    ("PRODUCTOS", "PRESENTACION"): ["CAJA", "UNIDAD", "BLISTER", "FRASCO", "BOLSA", "AMPOLLA", "TUBO", "DISPLAY"],
    ("PRODUCTOS", "UNIDAD"): ["UNIDAD", "CAJA", "FRASCO", "BLISTER", "AMPOLLA", "BOLSA"],
    ("PRODUCTOS", "TEMPERATURA"): ["AMBIENTE", "REFRIGERADO", "CONGELADO", "CONTROLADA"],
    ("PRODUCTOS", "ESTADO"): ["ACTIVO", "INACTIVO"],
    ("DOCUMENTOS", "TIPO_DOCUMENTO"): [
        "COA",
        "BPM",
        "BPA",
        "REGISTRO_SANITARIO",
        "GUIA_COMPRA",
        "FACTURA_COMPRA",
        "GUIA_VENTA",
        "FACTURA_VENTA",
        "OTRO",
    ],
}

DATES = {("LOTES", "FECHA_FABRICACION"), ("LOTES", "FECHA_VENCIMIENTO"), ("INGRESOS", "FECHA"), ("EGRESOS", "FECHA")}
NUMBERS = {("CONFIG", "DIAS_ALERTA"), ("LOTES", "PRECIO_COMPRA"), ("LOTES", "PRECIO_VENTA"), ("INGRESOS", "CANTIDAD"), ("EGRESOS", "CANTIDAD")}
IMAGES = {("CONFIG", "LOGO")}
FILES = {("DOCUMENTOS", "ARCHIVO")}

OPTIONAL = {
    ("CONFIG", "LOGO"),
    ("PROVEEDORES", "DIRECCION"),
    ("PROVEEDORES", "TELEFONO"),
    ("PROVEEDORES", "EMAIL"),
    ("PROVEEDORES", "CONTACTO"),
    ("CLIENTES", "DIRECCION"),
    ("CLIENTES", "TELEFONO"),
    ("CLIENTES", "EMAIL"),
    ("PRODUCTOS", "LABORATORIO"),
    ("PRODUCTOS", "REGISTRO_SANITARIO"),
    ("INGRESOS", "GUIA"),
    ("INGRESOS", "FACTURA"),
    ("EGRESOS", "GUIA"),
    ("EGRESOS", "FACTURA"),
    ("DOCUMENTOS", "OBSERVACION"),
}


VIRTUAL_COLUMNS = [
    {
        "table": "LOTES",
        "column": "RELATED_INGRESOS",
        "type": "List",
        "app_formula": 'REF_ROWS("INGRESOS", "ID_LOTE")',
        "purpose": "Trazabilidad: ingresos relacionados al lote.",
    },
    {
        "table": "LOTES",
        "column": "RELATED_EGRESOS",
        "type": "List",
        "app_formula": 'REF_ROWS("EGRESOS", "ID_LOTE")',
        "purpose": "Trazabilidad: egresos relacionados al lote.",
    },
    {
        "table": "LOTES",
        "column": "RELATED_DOCUMENTOS",
        "type": "List",
        "app_formula": 'REF_ROWS("DOCUMENTOS", "ID_LOTE")',
        "purpose": "Trazabilidad: documentos relacionados al lote.",
    },
    {
        "table": "LOTES",
        "column": "TRAZABILIDAD_RESUMEN",
        "type": "LongText",
        "app_formula": 'CONCATENATE("Producto: ", [ID_PRODUCTO].[NOMBRE], " | Lote: ", [LOTE], " | Stock: ", [STOCK_ACTUAL], " | Vence: ", TEXT([FECHA_VENCIMIENTO]))',
        "purpose": "Resumen visible para inspeccion.",
    },
    {
        "table": "LOTES",
        "column": "TOTAL_INGRESOS",
        "type": "Number",
        "app_formula": 'SUM(SELECT(INGRESOS[CANTIDAD], [ID_LOTE] = [_THISROW].[ID_LOTE]))',
        "purpose": "Total ingresado por lote.",
    },
    {
        "table": "LOTES",
        "column": "TOTAL_EGRESOS",
        "type": "Number",
        "app_formula": 'SUM(SELECT(EGRESOS[CANTIDAD], [ID_LOTE] = [_THISROW].[ID_LOTE]))',
        "purpose": "Total egresado por lote.",
    },
    {
        "table": "LOTES",
        "column": "STOCK_ACTUAL",
        "type": "Number",
        "app_formula": "[TOTAL_INGRESOS] - [TOTAL_EGRESOS]",
        "purpose": "Stock por lote.",
    },
    {
        "table": "LOTES",
        "column": "DIAS_VENCIMIENTO",
        "type": "Number",
        "app_formula": "[FECHA_VENCIMIENTO] - TODAY()",
        "purpose": "Dias restantes al vencimiento.",
    },
    {
        "table": "LOTES",
        "column": "ALERTA_VENCIMIENTO",
        "type": "Enum",
        "app_formula": 'IFS([STOCK_ACTUAL] <= 0, "SIN STOCK", [FECHA_VENCIMIENTO] < TODAY(), "VENCIDO", [DIAS_VENCIMIENTO] <= 30, "VENCE 30 DIAS", [DIAS_VENCIMIENTO] <= 60, "VENCE 60 DIAS", [DIAS_VENCIMIENTO] <= 90, "VENCE 90 DIAS", TRUE, "VIGENTE")',
        "purpose": "Alerta sanitaria.",
    },
    {
        "table": "LOTES",
        "column": "PRIMER_INGRESO",
        "type": "Date",
        "app_formula": 'MIN(SELECT(INGRESOS[FECHA], [ID_LOTE] = [_THISROW].[ID_LOTE]))',
        "purpose": "FIFO para desempate.",
    },
    {
        "table": "LOTES",
        "column": "ORDEN_FEFO_FIFO",
        "type": "Number",
        "app_formula": 'COUNT(SELECT(LOTES[ID_LOTE], AND([ID_PRODUCTO] = [_THISROW].[ID_PRODUCTO], [STOCK_ACTUAL] > 0, OR([FECHA_VENCIMIENTO] < [_THISROW].[FECHA_VENCIMIENTO], AND([FECHA_VENCIMIENTO] = [_THISROW].[FECHA_VENCIMIENTO], [PRIMER_INGRESO] < [_THISROW].[PRIMER_INGRESO]))))) + 1',
        "purpose": "Prioridad FEFO/FIFO.",
    },
    {
        "table": "LOTES",
        "column": "ES_LOTE_FEFO",
        "type": "Yes/No",
        "app_formula": "[ORDEN_FEFO_FIFO] = 1",
        "purpose": "Marca lote prioritario.",
    },
    {
        "table": "EGRESOS",
        "column": "STOCK_DISPONIBLE_LOTE",
        "type": "Number",
        "app_formula": "[ID_LOTE].[STOCK_ACTUAL]",
        "purpose": "Stock disponible para validar salida.",
    },
    {
        "table": "EGRESOS",
        "column": "ES_FEFO_VALIDO",
        "type": "Yes/No",
        "app_formula": '[ID_LOTE] = ANY(ORDERBY(SELECT(LOTES[ID_LOTE], AND([ID_PRODUCTO] = [_THISROW].[ID_PRODUCTO], [STOCK_ACTUAL] > 0)), [FECHA_VENCIMIENTO], TRUE, [PRIMER_INGRESO], TRUE))',
        "purpose": "Valida lote FEFO/FIFO.",
    },
]

SLICES = [
    {
        "name": "STOCK_DISPONIBLE",
        "source_table": "LOTES",
        "row_filter": "[STOCK_ACTUAL] > 0",
        "columns": "ID_LOTE,ID_PRODUCTO,LOTE,FECHA_VENCIMIENTO,STOCK_ACTUAL,ALERTA_VENCIMIENTO,ES_LOTE_FEFO",
    },
    {
        "name": "LOTES_ALERTA_DIGEMID",
        "source_table": "LOTES",
        "row_filter": 'AND([STOCK_ACTUAL] > 0, IN([ALERTA_VENCIMIENTO], LIST("VENCIDO", "VENCE 30 DIAS", "VENCE 60 DIAS", "VENCE 90 DIAS")))',
        "columns": "ID_LOTE,ID_PRODUCTO,LOTE,FECHA_VENCIMIENTO,STOCK_ACTUAL,ALERTA_VENCIMIENTO",
    },
    {
        "name": "TRAZABILIDAD_LOTES",
        "source_table": "LOTES",
        "row_filter": "TRUE",
        "columns": "ID_LOTE,ID_PRODUCTO,LOTE,TRAZABILIDAD_RESUMEN,RELATED_INGRESOS,RELATED_EGRESOS,RELATED_DOCUMENTOS",
    },
]

SECURITY_FILTERS = [
    {"table": "CONFIG", "filter": "TRUE"},
    {"table": "PROVEEDORES", "filter": "TRUE"},
    {"table": "CLIENTES", "filter": "TRUE"},
    {"table": "PRODUCTOS", "filter": "TRUE"},
    {"table": "LOTES", "filter": "TRUE"},
    {"table": "INGRESOS", "filter": "TRUE"},
    {"table": "EGRESOS", "filter": "TRUE"},
    {"table": "DOCUMENTOS", "filter": "TRUE"},
]

VIEWS = [
    {"name": "Dashboard", "type": "Dashboard", "data": "LOTES_ALERTA_DIGEMID,STOCK_DISPONIBLE,INGRESOS,EGRESOS", "position": "Primary", "icon": "dashboard"},
    {"name": "Ingresos", "type": "Table", "data": "INGRESOS", "position": "Primary", "icon": "download"},
    {"name": "Egresos", "type": "Table", "data": "EGRESOS", "position": "Primary", "icon": "upload"},
    {"name": "Stock", "type": "Table", "data": "STOCK_DISPONIBLE", "position": "Primary", "icon": "inventory"},
    {"name": "Trazabilidad", "type": "Detail", "data": "TRAZABILIDAD_LOTES", "position": "Primary", "icon": "search"},
    {"name": "Clientes", "type": "Deck", "data": "CLIENTES", "position": "Primary", "icon": "people"},
    {"name": "Proveedores", "type": "Deck", "data": "PROVEEDORES", "position": "Primary", "icon": "local_shipping"},
    {"name": "Configuración", "type": "Detail", "data": "CONFIG", "position": "Primary", "icon": "settings"},
    {"name": "Productos", "type": "Deck", "data": "PRODUCTOS", "position": "Menu", "icon": "medication"},
    {"name": "Lotes", "type": "Table", "data": "LOTES", "position": "Menu", "icon": "category"},
    {"name": "Documentos", "type": "Gallery", "data": "DOCUMENTOS", "position": "Menu", "icon": "attach_file"},
]

ACTIONS = [
    {"table": "LOTES", "name": "Ver trazabilidad del lote", "type": "App: go to another view within this app", "target": 'LINKTOROW([ID_LOTE], "Trazabilidad_Detail")'},
    {"table": "PRODUCTOS", "name": "Ver lotes del producto", "type": "App: go to another view within this app", "target": 'LINKTOFILTEREDVIEW("Stock", [ID_PRODUCTO] = [_THISROW].[ID_PRODUCTO])'},
    {"table": "LOTES", "name": "Registrar egreso de lote", "type": "App: go to another view within this app", "target": 'LINKTOFORM("EGRESOS_Form", "ID_PRODUCTO", [ID_PRODUCTO], "ID_LOTE", [ID_LOTE])'},
    {"table": "LOTES", "name": "Registrar documento de lote", "type": "App: go to another view within this app", "target": 'LINKTOFORM("DOCUMENTOS_Form", "ID_LOTE", [ID_LOTE])'},
    {"table": "INGRESOS", "name": "Abrir lote", "type": "App: go to another view within this app", "target": 'LINKTOROW([ID_LOTE], "LOTES_Detail")'},
    {"table": "EGRESOS", "name": "Abrir cliente", "type": "App: go to another view within this app", "target": 'LINKTOROW([ID_CLIENTE], "CLIENTES_Detail")'},
]

BOTS = [
    {"name": "Alerta lotes vencidos", "event": "Scheduled daily", "table": "LOTES", "condition": 'SELECT(LOTES[ID_LOTE], AND([STOCK_ACTUAL] > 0, [FECHA_VENCIMIENTO] < TODAY()))', "task": "Email"},
    {"name": "Alerta lotes por vencer 30 dias", "event": "Scheduled daily", "table": "LOTES", "condition": 'SELECT(LOTES[ID_LOTE], AND([STOCK_ACTUAL] > 0, [DIAS_VENCIMIENTO] >= 0, [DIAS_VENCIMIENTO] <= 30))', "task": "Email"},
    {"name": "Alerta egreso registrado", "event": "Data change adds only", "table": "EGRESOS", "condition": "TRUE", "task": "Notification"},
    {"name": "Alerta documento faltante", "event": "Scheduled weekly", "table": "LOTES", "condition": 'SELECT(LOTES[ID_LOTE], AND([STOCK_ACTUAL] > 0, COUNT([RELATED_DOCUMENTOS]) = 0))', "task": "Email"},
]

FORMAT_RULES = [
    {"name": "Vencido", "table": "LOTES", "condition": '[ALERTA_VENCIMIENTO] = "VENCIDO"', "color": "#B42318", "icon": "warning"},
    {"name": "Vence 30 dias", "table": "LOTES", "condition": '[ALERTA_VENCIMIENTO] = "VENCE 30 DIAS"', "color": "#DC6803", "icon": "schedule"},
    {"name": "Vence 60 dias", "table": "LOTES", "condition": '[ALERTA_VENCIMIENTO] = "VENCE 60 DIAS"', "color": "#B54708", "icon": "schedule"},
    {"name": "Vence 90 dias", "table": "LOTES", "condition": '[ALERTA_VENCIMIENTO] = "VENCE 90 DIAS"', "color": "#027A48", "icon": "check_circle"},
    {"name": "Sin stock", "table": "LOTES", "condition": "[STOCK_ACTUAL] <= 0", "color": "#667085", "icon": "block"},
    {"name": "FEFO prioritario", "table": "LOTES", "condition": "[ES_LOTE_FEFO] = TRUE", "color": "#0F766E", "icon": "priority_high"},
]

UX = {
    "starting_view": "Dashboard",
    "primary_navigation": ["Dashboard", "Ingresos", "Egresos", "Stock", "Trazabilidad", "Clientes", "Proveedores", "Configuración"],
    "menu_navigation": ["Productos", "Lotes", "Documentos"],
    "dashboard_interactive_mode": True,
    "search_enabled_for": ["LOTES", "PRODUCTOS", "CLIENTES", "PROVEEDORES"],
}

BRANDING = {
    "app_name": "Droguería - Trazabilidad DIGEMID",
    "logo_source": "CONFIG[LOGO]",
    "theme": "Light",
    "primary_color": "#0F766E",
    "background_color": "#FFFFFF",
    "header_logo": True,
}


def appsheet_type(table: str, column: str) -> str:
    key = (table, column)
    if column == KEYS[table]:
        return "Text"
    if key in REFS:
        return "Ref"
    if key in ENUMS:
        return "Enum"
    if key in DATES:
        return "Date"
    if key in NUMBERS:
        return "Number"
    if key in IMAGES:
        return "Image"
    if key in FILES:
        return "File"
    if "EMAIL" in column:
        return "Email"
    if "TELEFONO" in column:
        return "Phone"
    return "Text"


def initial_value(table: str, column: str) -> str:
    if column == KEYS[table]:
        return "UNIQUEID()"
    if column == "FECHA":
        return "TODAY()"
    if column == "USUARIO":
        return "USEREMAIL()"
    if column == "ESTADO":
        return '"ACTIVO"'
    if (table, column) == ("CONFIG", "DIAS_ALERTA"):
        return "90"
    return ""


def valid_if(table: str, column: str) -> str:
    key = (table, column)
    if key in ENUMS:
        return "LIST(" + ",".join(f'"{v}"' for v in ENUMS[key]) + ")"
    if (table, column) == ("INGRESOS", "ID_LOTE"):
        return "SELECT(LOTES[ID_LOTE], [ID_PRODUCTO] = [_THISROW].[ID_PRODUCTO])"
    if (table, column) == ("EGRESOS", "ID_LOTE"):
        return "LIST(ANY(ORDERBY(SELECT(LOTES[ID_LOTE], AND([ID_PRODUCTO] = [_THISROW].[ID_PRODUCTO], [STOCK_ACTUAL] > 0)), [FECHA_VENCIMIENTO], TRUE, [PRIMER_INGRESO], TRUE)))"
    if (table, column) == ("EGRESOS", "CANTIDAD"):
        return "AND([_THIS] > 0, [_THIS] <= [ID_LOTE].[STOCK_ACTUAL])"
    if (table, column) == ("INGRESOS", "CANTIDAD"):
        return "[_THIS] > 0"
    if (table, column) == ("LOTES", "FECHA_VENCIMIENTO"):
        return "[FECHA_VENCIMIENTO] >= [FECHA_FABRICACION]"
    if (table, column) == ("PRODUCTOS", "CODIGO"):
        return "COUNT(SELECT(PRODUCTOS[ID_PRODUCTO], AND([CODIGO] = [_THISROW].[CODIGO], [ID_PRODUCTO] <> [_THISROW].[ID_PRODUCTO]))) = 0"
    return ""


def base_column_config() -> list[dict]:
    rows = []
    for table, columns in TABLES.items():
        for column in columns:
            key = (table, column)
            rows.append(
                {
                    "table": table,
                    "column": column,
                    "type": appsheet_type(table, column),
                    "key": column == KEYS[table],
                    "label": column == LABELS[table],
                    "ref_table": REFS.get(key, ""),
                    "enum_values": ENUMS.get(key, []),
                    "required": key not in OPTIONAL,
                    "editable": column != KEYS[table],
                    "initial_value": initial_value(table, column),
                    "valid_if": valid_if(table, column),
                    "app_formula": "",
                    "show": True,
                }
            )
    for item in VIRTUAL_COLUMNS:
        rows.append(
            {
                "table": item["table"],
                "column": item["column"],
                "type": item["type"],
                "key": False,
                "label": False,
                "ref_table": "",
                "enum_values": [],
                "required": False,
                "editable": False,
                "initial_value": "",
                "valid_if": "",
                "app_formula": item["app_formula"],
                "show": True,
            }
        )
    return rows


def write_csv(path: Path, rows: list[dict], headers: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, data) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def build_excel() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, columns in TABLES.items():
        ws = wb.create_sheet(name)
        ws.append(columns)
    wb.save(XLSX)


def build_package() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_excel()
    columns = base_column_config()

    write_json(
        OUT_DIR / "appsheet_app_spec.json",
        {
            "app": {
                "name": "Droguería - Trazabilidad DIGEMID",
                "platform": "AppSheet",
                "data_source": "Google Sheets",
                "production_ready": True,
            },
            "tables": TABLES,
            "columns": columns,
            "virtual_columns": VIRTUAL_COLUMNS,
            "slices": SLICES,
            "security_filters": SECURITY_FILTERS,
            "views": VIEWS,
            "actions": ACTIONS,
            "bots": BOTS,
            "format_rules": FORMAT_RULES,
            "ux": UX,
            "branding": BRANDING,
        },
    )
    write_csv(OUT_DIR / "01_tables.csv", [{"table": t, "key": KEYS[t], "label": LABELS[t], "columns": ",".join(c)} for t, c in TABLES.items()], ["table", "key", "label", "columns"])
    write_csv(OUT_DIR / "02_columns.csv", columns, ["table", "column", "type", "key", "label", "ref_table", "enum_values", "required", "editable", "initial_value", "valid_if", "app_formula", "show"])
    write_csv(OUT_DIR / "03_virtual_columns.csv", VIRTUAL_COLUMNS, ["table", "column", "type", "app_formula", "purpose"])
    write_csv(OUT_DIR / "04_slices.csv", SLICES, ["name", "source_table", "row_filter", "columns"])
    write_csv(OUT_DIR / "05_security_filters.csv", SECURITY_FILTERS, ["table", "filter"])
    write_csv(OUT_DIR / "06_views.csv", VIEWS, ["name", "type", "data", "position", "icon"])
    write_csv(OUT_DIR / "07_actions.csv", ACTIONS, ["table", "name", "type", "target"])
    write_csv(OUT_DIR / "08_bots.csv", BOTS, ["name", "event", "table", "condition", "task"])
    write_csv(OUT_DIR / "09_format_rules.csv", FORMAT_RULES, ["name", "table", "condition", "color", "icon"])
    write_json(OUT_DIR / "10_ux_branding.json", {"ux": UX, "branding": BRANDING})
    write_json(OUT_DIR / "11_navigation.json", {"primary": UX["primary_navigation"], "menu": UX["menu_navigation"]})
    (OUT_DIR / "12_expressions.md").write_text(expressions_doc(), encoding="utf-8")
    IMPLEMENTACION.write_text(implementation_doc(), encoding="utf-8")
    verify()


def expressions_doc() -> str:
    lines = ["# Expresiones AppSheet Producción", ""]
    for item in VIRTUAL_COLUMNS:
        lines += [f"## {item['table']}[{item['column']}]", "", "```appsheet", item["app_formula"], "```", ""]
    lines += ["# Valid If críticos", ""]
    for row in base_column_config():
        if row["valid_if"]:
            lines += [f"## {row['table']}[{row['column']}]", "", "```appsheet", row["valid_if"], "```", ""]
    return "\n".join(lines)


def implementation_doc() -> str:
    return "\n".join(
        [
            "# IMPLEMENTACION_FINAL",
            "",
            "1. Subir `DIGEMID_Drogueria_AppSheet.xlsx` a Google Drive.",
            "2. Abrir el archivo con Google Sheets.",
            "3. Renombrar el Google Sheet como `BD_DIGEMID_DROGUERIA_PRODUCCION`.",
            "4. Entrar a `appsheet.com`.",
            "5. Clic en `Create`.",
            "6. Clic en `App`.",
            "7. Clic en `Start with existing data`.",
            "8. Seleccionar `BD_DIGEMID_DROGUERIA_PRODUCCION`.",
            "9. Crear la app.",
            "10. En `Data > Tables`, confirmar estas tablas: `CONFIG`, `PROVEEDORES`, `CLIENTES`, `PRODUCTOS`, `LOTES`, `INGRESOS`, `EGRESOS`, `DOCUMENTOS`.",
            "11. En `Data > Columns`, aplicar la configuración de `02_columns.csv`.",
            "12. En `Data > Columns`, crear las columnas virtuales de `03_virtual_columns.csv`.",
            "13. En `Data > Slices`, crear los slices de `04_slices.csv`.",
            "14. En `Security > Security Filters`, aplicar `05_security_filters.csv`.",
            "15. En `UX > Views`, crear las vistas de `06_views.csv`.",
            "16. En `Behavior > Actions`, crear las acciones de `07_actions.csv`.",
            "17. En `Automation > Bots`, crear los bots de `08_bots.csv`.",
            "18. En `UX > Format Rules`, crear las reglas de `09_format_rules.csv`.",
            "19. En `UX > Brand`, aplicar `10_ux_branding.json`.",
            "20. En `UX > Views`, dejar como navegación principal: Dashboard, Ingresos, Egresos, Stock, Trazabilidad, Clientes, Proveedores, Configuración.",
            "21. Clic en `Save`.",
            "22. Ejecutar prueba: crear producto, proveedor, lote, ingreso, cliente y egreso.",
            "23. Verificar que `STOCK_ACTUAL` baja después del egreso.",
            "24. Verificar que AppSheet bloquea egreso mayor al stock.",
            "25. Verificar que AppSheet solo permite el lote FEFO/FIFO.",
            "26. Verificar que Trazabilidad muestra ingresos, egresos y documentos del lote.",
            "27. Clic en `Manage > Deploy`.",
            "28. Resolver únicamente advertencias obligatorias de AppSheet.",
            "29. Clic en `Move app to deployed state`.",
        ]
    )


def verify() -> None:
    wb = load_workbook(XLSX, data_only=False)
    assert wb.sheetnames == list(TABLES)
    assert all(wb[s].max_row == 1 for s in TABLES)
    assert sum(len(wb[s].merged_cells.ranges) for s in wb.sheetnames) == 0
    required = [
        "appsheet_app_spec.json",
        "01_tables.csv",
        "02_columns.csv",
        "03_virtual_columns.csv",
        "04_slices.csv",
        "05_security_filters.csv",
        "06_views.csv",
        "07_actions.csv",
        "08_bots.csv",
        "09_format_rules.csv",
        "10_ux_branding.json",
        "11_navigation.json",
        "12_expressions.md",
        "IMPLEMENTACION_FINAL.md",
    ]
    for name in required:
        assert (OUT_DIR / name).exists(), name


if __name__ == "__main__":
    build_package()
    print(OUT_DIR)
